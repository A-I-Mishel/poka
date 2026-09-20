"""Teaching containment: window-only input for teaching turns.

A teaching turn must never seed the full-deck text hint: the teaching
stage's window hint carries only the current slides, and the model must
not see slides it was not shown. Non-teaching turns keep the full text
seed so document Q&A keeps working. (Stubbed agent, real vault.)
"""

import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))


def _ctx(tmp_path, monkeypatch, uid="teach-user"):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("PLUTO_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("PLUTO_USER_ID", uid)
    monkeypatch.delenv("PLUTO_AUTH_MODE", raising=False)
    from backend.deps import UserContext
    from services.files import FileStore
    from services.storage import UserStore

    return UserContext(user_id=uid, user_store=UserStore(uid),
                       file_store=FileStore(uid), limit_key=uid, source="env")


def _five_slide_deck():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[6]  # blank: shapes fully ours
    bodies = [
        "Alpha concept: search trees expand nodes level by level, "
        "tracking visited states carefully to avoid repeat work and "
        "wasted effort during traversal of large problem spaces.",
        "Beta concept: heuristics estimate distance to the goal quickly, "
        "guiding expansion toward promising regions first while keeping "
        "the search focused and efficient throughout the whole process.",
        "Gamma concept: A star balances cost so far and estimate together, "
        "adding both terms into one evaluation function that guarantees "
        "optimal paths when the heuristic never overestimates true cost.",
        "Delta concept: greedy search picks the best successor only, "
        "ignoring path history entirely, which runs fast but risks missing "
        "better routes hidden behind temporarily worse-looking first moves.",
        "Epsilon concept: plateaus trap local search without progress, "
        "since flat regions offer no gradient signal at all, forcing random "
        "restarts or sideways moves to escape the featureless dead zones.",
    ]
    for i, body in enumerate(bodies, start=1):
        slide = prs.slides.add_slide(layout)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5),
                                       Inches(9), Inches(2))
        box.text_frame.text = f"Slide {i}. {body}"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _upload(ctx, data=b"deck notes body " + b"y" * 50):
    return ctx.file_store.save_upload(data, "lecture.pptx")


def _fake_agent(monkeypatch, seen):
    import agent as agent_mod

    def _answer(user_input, history=None, **kwargs):
        seen["input"] = str(user_input)
        return {"output": "ok", "active_tier": "Fake",
                "task_type": "simple", "tools_used": [], "sources": []}

    monkeypatch.setattr(agent_mod, "answer_with_fallback", _answer)


def test_teaching_turn_skips_full_deck_seed(tmp_path, monkeypatch):
    from backend.chatflow import run_chat

    ctx = _ctx(tmp_path, monkeypatch, "teach-seed")
    meta = _upload(ctx, _five_slide_deck())
    seen = {}
    _fake_agent(monkeypatch, seen)
    run_chat(ctx, "teach me this lecture slide by slide",
             upload_ids=[meta.id])
    sent = seen["input"]
    # Window hint present (current slides only)...
    assert "Alpha concept" in sent
    assert "[slide 1]" in sent
    # ...tool pointer present so read_document keeps working...
    assert meta.id in sent
    # ...but slides outside the first window were never shown.
    assert "Epsilon concept" not in sent
    assert "Delta concept" not in sent


def test_nonteaching_turn_keeps_full_deck_seed(tmp_path, monkeypatch):
    from backend.chatflow import run_chat

    ctx = _ctx(tmp_path, monkeypatch, "teach-noseed")
    meta = _upload(ctx, _five_slide_deck())
    seen = {}
    _fake_agent(monkeypatch, seen)
    run_chat(ctx, "summarize this document please", upload_ids=[meta.id])
    sent = seen["input"]
    assert "Alpha concept" in sent
    assert "Epsilon concept" in sent


def _title_only_deck():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[6]
    for i in ("Alpha", "Beta", "Gamma"):
        slide = prs.slides.add_slide(layout)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5),
                                       Inches(9), Inches(1))
        box.text_frame.text = i
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_full_window_carries_no_refetch_note(tmp_path, monkeypatch):
    from backend.chatflow import run_chat

    ctx = _ctx(tmp_path, monkeypatch, "teach-norefetch")
    meta = _upload(ctx, _five_slide_deck())
    seen = {}
    _fake_agent(monkeypatch, seen)
    run_chat(ctx, "teach me this lecture slide by slide",
             upload_ids=[meta.id])
    sent = seen["input"]
    assert "do not re-fetch" in sent
    # Tool pointer stays so diagrams/overflow still work.
    assert meta.id in sent


def test_thin_window_keeps_fetch_behavior(tmp_path, monkeypatch):
    from backend.chatflow import run_chat

    ctx = _ctx(tmp_path, monkeypatch, "teach-thin")
    meta = _upload(ctx, _title_only_deck())
    seen = {}
    _fake_agent(monkeypatch, seen)
    run_chat(ctx, "teach me this lecture slide by slide",
             upload_ids=[meta.id])
    sent = seen["input"]
    assert "title-only" in sent
    assert "do not re-fetch" not in sent


def test_system_prompt_softened_fetch_order():
    import agent.prompts as prompts_mod

    assert "without re-fetching" in prompts_mod.SYSTEM_PROMPT
    # Pinned concept/format tokens survive the rewording.
    assert "## Concept:" in prompts_mod.SYSTEM_PROMPT
