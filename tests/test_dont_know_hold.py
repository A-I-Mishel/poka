"""Don't-know hold: failed recall reteaches, never advances.

Screenshot regression: "NOPE I CANT THINK" was ignored and slide 11
was taught over it. Now the cursor holds on the same window with a
reteach directive; a repeat hold escalates (prerequisite + offer to
park it). Attempted answers keep evaluate-then-advance.
"""

import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))


def _pptx_bytes(slides):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for texts in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for i, txt in enumerate(texts):
            tx = slide.shapes.add_textbox(
                Inches(0.5), Inches(0.5 + i), Inches(5), Inches(1)
            )
            tx.text_frame.text = txt
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _ctx(tmp_path, monkeypatch, uid="teach-hold"):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("PLUTO_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("PLUTO_USER_ID", uid)
    monkeypatch.delenv("PLUTO_AUTH_MODE", raising=False)
    from backend.deps import UserContext
    from services.files import FileStore
    from services.storage import UserStore

    return UserContext(user_id=uid, user_store=UserStore(uid),
                       file_store=FileStore(uid), limit_key=uid, source="env")


def _lesson(n):
    return (f"📘 FILE: Lecture_01.pptx\nSlides: {n}-{n}\n## Concept: G\n"
            f"**Source:** [slide {n}]\nWhat is S{n}?")


def test_detector_positives():
    from backend.chatflow import _is_dont_know

    hist = [{"role": "assistant", "content": _lesson(10)}]
    for text in ("NOPE I CANT THINK", "i dont know", "I don't know",
                 "idk", "no idea", "no clue", "skip this one",
                 "nope, cant think of anything"):
        assert _is_dont_know(text, hist) is True, text


def test_detector_negatives():
    from backend.chatflow import _is_dont_know

    hist = [{"role": "assistant", "content": _lesson(10)}]
    for text in ("A vertex is a node", "It counts connected edges",
                 "not sure, is it X?", "Next", "ok",
                 "created from vertices and edges"):
        assert _is_dont_know(text, hist) is False, text
    # No teaching question pending: not a hold either.
    assert _is_dont_know(
        "i dont know",
        [{"role": "assistant", "content": "hello there"}]) is False


def test_hold_reserves_same_window(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch)
    bodies = [[f"Slide {i} graph vertex edge"] for i in range(1, 12)]
    m1 = ctx.file_store.save_upload(_pptx_bytes(bodies), "Lecture_01.pptx")
    hist = [
        {"role": "user", "content": "teach me", "attachments": [
            {"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]},
        {"role": "assistant", "content": _lesson(10)},
    ]
    send, _, clarify = _apply_teaching_session(
        ctx, "NOPE I CANT THINK", hist, [], [], "NOPE I CANT THINK")
    assert clarify is None
    assert "Slide 10" in send
    assert "Slide 11" not in send
    assert "Do NOT advance" in send
    # Softer reteach: shape restated, learner wording reused, fresh metaphor.
    assert "Imagine" in send
    assert "metaphor" in send.lower()


def test_repeat_hold_escalates(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-hold2")
    bodies = [[f"Slide {i} graph vertex edge"] for i in range(1, 12)]
    m1 = ctx.file_store.save_upload(_pptx_bytes(bodies), "Lecture_01.pptx")
    hist = [
        {"role": "user", "content": "teach me", "attachments": [
            {"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]},
        {"role": "assistant", "content": _lesson(10)},
        {"role": "user", "content": "i dont know"},
        {"role": "assistant", "content": _lesson(10)},
    ]
    send, _, clarify = _apply_teaching_session(
        ctx, "still cant think", hist, [], [], "still cant think")
    assert clarify is None
    assert "Slide 10" in send
    assert "park it" in send


def test_attempt_still_advances(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-hold3")
    bodies = [[f"Slide {i} graph vertex edge"] for i in range(1, 12)]
    m1 = ctx.file_store.save_upload(_pptx_bytes(bodies), "Lecture_01.pptx")
    hist = [
        {"role": "user", "content": "teach me", "attachments": [
            {"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]},
        {"role": "assistant", "content": _lesson(10)},
    ]
    send, _, clarify = _apply_teaching_session(
        ctx, "It counts connected edges", hist, [], [],
        "It counts connected edges")
    assert clarify is None
    assert "Slide 11" in send
    assert "evaluate" in send.lower()
    # Wrong answers get the softer repair too (then advance, per design).
    assert "metaphor" in send.lower()
