"""PPTX embedded pictures: alt-text + OCR ladder (stubbed, zero quota).

Synthetic decks built in-test (PIL PNG bytes, no fixtures on disk).
Covers the shared helper plus all three readers (teaching blocks,
document file, KB flat text).
"""

import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))


def _png_bytes(color=(255, 255, 255)):
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (32, 32), color).save(buf, format="PNG")
    return buf.getvalue()


def _deck(slides):
    """slides: list of (texts, pictures) where pictures is a list of
    alt-or-None per picture. Returns pptx bytes."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for texts, pics in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for i, txt in enumerate(texts):
            tx = slide.shapes.add_textbox(
                Inches(0.5), Inches(0.5 + i), Inches(5), Inches(1))
            tx.text_frame.text = txt
        for alt in pics:
            pic = slide.shapes.add_picture(
                io.BytesIO(_png_bytes()), Inches(0.5), Inches(3),
                Inches(2), Inches(2))
            if alt is not None:
                pic._element.nvPicPr.cNvPr.set("descr", alt)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_alt_text_preferred_no_ocr_called(monkeypatch):
    from services import pptx_images as pi

    called = []
    monkeypatch.setattr(pi, "ocr_picture_on_device",
                        lambda blob: called.append(blob) or "SHOULD-NOT-RUN")

    def _boom(blob):
        raise AssertionError("vision OCR must not run when alt text exists")

    lines, used = pi.picture_lines_for_slide(
        [("Osmosis process diagram", b"fake-bytes")], 1, _boom)
    assert used == 1
    assert lines == ["[image 1: Osmosis process diagram]"]
    assert called == []


def test_auto_name_is_not_alt_text():
    from services.pptx_images import picture_alt_text

    class FakeEl:
        pass

    class FakeShape:
        name = "Picture 7"
        _element = FakeEl()

    assert picture_alt_text(FakeShape()) == ""


def test_builder_default_descr_is_not_alt_text():
    # python-pptx add_picture prefills descr with the file name
    # ("image.png") — builder noise, not author-written alt text.
    from services.pptx_images import picture_alt_text

    class FakeCnv:
        def get(self, key):
            return "image.png" if key == "descr" else ""

    class FakeNv:
        cNvPr = FakeCnv()

    class FakeEl:
        nvPicPr = FakeNv()

    class FakeShape:
        name = "Picture 3"
        _element = FakeEl()

    assert picture_alt_text(FakeShape()) == ""


def test_descr_hop_failure_falls_back_to_name():
    from services.pptx_images import picture_alt_text

    class FakeShape:
        name = "My custom diagram"

    assert picture_alt_text(FakeShape()) == "My custom diagram"


def test_ocr_ladder_on_device_then_vision(monkeypatch):
    from services import pptx_images as pi

    monkeypatch.setattr(pi, "ocr_picture_on_device", lambda blob: "OCR WORDS")
    lines, _ = pi.picture_lines_for_slide([( "", b"b")], 2, None)
    assert lines == ["[image 2 (OCR):\nOCR WORDS]"]

    monkeypatch.setattr(pi, "ocr_picture_on_device", lambda blob: "")
    lines, _ = pi.picture_lines_for_slide(
        [( "", b"b")], 3, lambda blob: "VISION WORDS")
    assert lines == ["[image 3 (vision-OCR):\nVISION WORDS]"]

    lines, used = pi.picture_lines_for_slide([( "", b"b")], 4, None)
    assert lines == [] and used == 1  # skipped silently, counter moves


def test_bounds_deck_and_slide(monkeypatch):
    from services import pptx_images as pi

    monkeypatch.setattr(pi, "ocr_picture_on_device", lambda blob: "")
    pics = [("Alt %d" % i, None) for i in range(8)]
    lines, used = pi.picture_lines_for_slide(pics, 1, None)
    assert used == 8 and len(lines) == 8  # per-call: caller caps via iter
    # Deck-level cap lives in iter_deck_pictures (tested below).
    assert pi.MAX_PPTX_IMAGES_PER_DECK == 6
    assert pi.MAX_PPTX_IMAGES_PER_SLIDE == 2


def test_iter_deck_pictures_bounds_and_order():
    from pptx import Presentation

    from services.pptx_images import iter_deck_pictures
    data = _deck([
        (["S1"], ["A1", "A2", "A3"]),
        (["S2"], ["B1", "B2"]),
        (["S3"], ["C1", "C2"]),
        (["S4"], ["D1", "D2"]),
    ])
    prs = Presentation(io.BytesIO(data))
    got = list(iter_deck_pictures(prs))
    # Per-slide cap 2 (A3 skipped), then deck cap 6 (slide 4 cut off).
    assert [g[0] for g in got] == [1, 1, 2, 2, 3, 3]
    assert len(got) == 6
    assert [g[1] for g in got][:2] == ["A1", "A2"]
    assert all(g[2] for g in got)  # blobs present


def test_teaching_blocks_include_pictures():
    from backend.teach import _extract_pptx_blocks
    import tempfile

    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as fh:
        fh.write(_make_deck_file())
        path = fh.name
    try:
        blocks = _extract_pptx_blocks(path, "pptx")
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass
    assert blocks and "[image 1: Osmosis process diagram]" in blocks[0][1]
    assert "Slide text here" in blocks[0][1]


def test_document_reader_includes_pictures(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    from tools.document_tool import _read_pptx_file

    p = tmp_path / "deck.pptx"
    p.write_bytes(_make_deck_file())
    out = _read_pptx_file(p)
    assert "[slide 1]" in out
    assert "[image 1: Osmosis process diagram]" in out


def test_kb_reader_includes_pictures_no_vision():
    from services import kb as kb_mod

    text, reason = kb_mod._pptx_text(_make_deck_file())
    assert reason == ""
    assert "[image 1: Osmosis process diagram]" in text


def _make_deck_file():
    return _deck([
        (["Slide text here"], ["Osmosis process diagram"]),
        (["Second slide"], [None]),
    ])
