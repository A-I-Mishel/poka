"""Teaching voice (humanization mission): 1-slide windows, demo grammar.

Hermetic: pure prompt-data logic, no network, no LLM calls.
"""

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))


def test_window_is_one_slide():
    from backend import teach as teach_mod

    assert teach_mod.TEACHING_WINDOW_SLIDES == 1
    blocks = [(1, "aaa"), (2, "bbb"), (3, "ccc")]
    window, start, end, truncated = teach_mod._select_teaching_window(blocks, 0)
    assert (start, end) == (1, 1)
    assert [n for n, _ in window] == [1]
    window2, s2, e2, _ = teach_mod._select_teaching_window(blocks, 1)
    assert (s2, e2) == (2, 2)
    assert [n for n, _ in window2] == [2]


def _demo_draft() -> str:
    return (
        "📘 FILE: Lecture_01.pptx\nSlides: 1-1\n"
        "Building on the course intro, here is the first real idea.\n"
        "## Concept: What is a Graph? (Slide 1)\n"
        "A graph is a set of vertices plus a set of edges that relate them.\n"
        "Think objects plus the relationships between them.\n"
        "Imagine:\n```\nA -------- B\n \\       /\n  \\     /\n   C\n```\n"
        "Here:\n- A, B, C = vertices\n- Lines = edges\n"
        "⭐ MUST MEMORIZE\n"
        "Vertex = object/node; Edge = connection/relationship.\n"
        "**Source:** [slide 1]\n"
        "**Recall**\nWhat is the difference between a vertex and an edge?\n"
        "Reply **continue** for the next concept."
    )


def test_demo_style_draft_passes_validation():
    from backend.teach import _validate_teaching_draft

    assert _validate_teaching_draft(_demo_draft(), 1, 1) == []


def test_demo_draft_has_no_robotic_template():
    draft = _demo_draft()
    assert "Exam trap" not in draft
    assert "Exam importance" not in draft
    assert draft.count("**Recall**") == 1
    assert "Imagine:" in draft and "Here:" in draft


def test_admin_compact_passes_without_recall():
    from backend.teach import _validate_teaching_draft

    admin = (
        "📘 FILE: Lecture_01.pptx\nSlides: 1-1\n"
        "### Administrative Information\n"
        "- Course: Graph Theory (CSE 0613-4125)\n- Credit: 3.00\n"
        "**Source:** [slide 1]\n"
        "Nothing technical here. Let's move on."
    )
    assert _validate_teaching_draft(admin, 1, 1) == []


def test_teaching_skips_reflection():
    from agent.reflection import should_reflect

    long_draft = "x" * 500
    teaching_input = "teach me\n\n[Teaching mode: warm human tutor]"
    assert should_reflect("research", long_draft, teaching_input, False) is False
    # Non-teaching research still reflects.
    assert should_reflect("research", long_draft, "plain question", False) is True


def test_fast_tiers_exclude_weakest_lane():
    import config

    fast = [n for n, _ in config.FAST_TIERS]
    synth = [n for n, _ in config.SYNTHESIS_TIERS]
    assert "Ollama 8B" not in fast
    assert "Ollama 8B" in synth  # deep-only offline tail kept
    assert set(fast) < set(synth)


def _teaching_input(extra: str = "") -> str:
    return (
        "teach me\n\n[Teaching mode: warm human tutor]\n"
        "[Verified content of 'L.pptx' slides 1-1 of 5:\n[slide 1]\nBody]\n"
        + extra
    )


def test_clean_window_unbinds_fetch_tools():
    from agent.toolrun import filter_tools_for_hint

    names = {getattr(t, "name", "") for t in filter_tools_for_hint(_teaching_input())}
    assert "read_document" not in names
    assert "read_pdf" not in names
    assert "read_pdf_page" not in names


def test_thin_window_keeps_fetch_tools():
    from agent.toolrun import filter_tools_for_hint

    hint = _teaching_input(
        "[Note: this window is title-only (12 chars of body text). "
        "Call read_pdf_page for pages 1-1 first.]")
    names = {getattr(t, "name", "") for t in filter_tools_for_hint(hint)}
    assert "read_pdf_page" in names


def test_truncated_window_keeps_fetch_tools():
    from agent.toolrun import filter_tools_for_hint

    hint = _teaching_input("[Note: window text truncated to fit context.]")
    names = {getattr(t, "name", "") for t in filter_tools_for_hint(hint)}
    assert "read_document" in names


def test_non_teaching_hints_unaffected():
    from agent.toolrun import filter_tools_for_hint

    names = {getattr(t, "name", "") for t in filter_tools_for_hint("read this document please")}
    assert "read_document" in names


def test_execution_drop_keeps_other_calls():
    from agent.toolrun import _drop_teaching_refetch_calls

    calls = [
        {"name": "read_document", "args": {"upload_id": "abc"}},
        {"name": "web_search", "args": {"query": "graphs"}},
    ]
    kept, dropped = _drop_teaching_refetch_calls(calls, _teaching_input())
    assert dropped == 1
    assert [c["name"] for c in kept] == ["web_search"]
    # Non-teaching input: nothing dropped.
    kept2, dropped2 = _drop_teaching_refetch_calls(calls, "plain question")
    assert dropped2 == 0 and len(kept2) == 2


def test_teaching_pointer_is_fetch_neutral():
    from backend.attachments import attachment_hint

    uid = "a" * 16
    neutral = attachment_hint("document", uid, "L.pptx", 1, 1, teaching=True)
    assert uid in neutral
    assert "To read it, call" not in neutral
    assert "do not re-fetch" in neutral
    classic = attachment_hint("document", uid, "L.pptx", 1, 1)
    assert "To read it, call read_document" in classic


def test_thin_turn_keeps_fetch_pointer(tmp_path, monkeypatch):
    """Thin/diagram windows keep the classic fetch pointer (diagrams stay reachable)."""
    from backend.chatflow import _apply_teaching_session
    from backend.deps import UserContext
    from services.files import FileStore
    from services.storage import UserStore

    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("PLUTO_DATA_DIR", str(tmp_path / "data"))
    ctx = UserContext(user_id="thin-ptr", user_store=UserStore("thin-ptr"),
                      file_store=FileStore("thin-ptr"), limit_key="thin-ptr", source="env")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    box.text_frame.text = "Diagrams"
    import io as _io

    buf = _io.BytesIO()
    prs.save(buf)
    meta = ctx.file_store.save_upload(buf.getvalue(), "Thin.pptx")
    atts = [{"id": meta.id, "kind": "document", "name": "Thin.pptx"}]
    send, _, _ = _apply_teaching_session(ctx, "teach me", [], atts, [], "teach me")
    assert "title-only" in send
    assert "To read it, call" in send
    assert "do not re-fetch" not in send


def test_doc_inline_html_renders_and_escapes():
    import html as _html

    from tools.make_tool import _inline_html

    # Contract: caller escapes first, then renders markers.
    out = _inline_html(_html.escape("A **bold** move with <script> and *em* plus `code`"))
    assert "<b>bold</b>" in out
    assert "<i>em</i>" in out
    assert "<code>code</code>" in out
    # Raw HTML is escaped by the caller; only our own tags are markup.
    assert "<script>" not in out
    assert "&lt;script&gt;" in out
