"""Stateless teaching session: ONE file + 3-slide window, fail-closed."""

import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))


def _pptx_bytes(slides):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for texts in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for i, txt in enumerate(texts):
            tx = slide.shapes.add_textbox(
                Inches(0.5), Inches(0.5 + i), Inches(5), Inches(1)
            )
            tx.text_frame.text = txt
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _ctx(tmp_path, monkeypatch, uid="teach-user"):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("PLUTO_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("PLUTO_USER_ID", uid)
    monkeypatch.delenv("PLUTO_AUTH_MODE", raising=False)
    from backend.deps import UserContext
    from services.files import FileStore
    from services.storage import UserStore

    return UserContext(user_id=uid, user_store=UserStore(uid),
                       file_store=FileStore(uid), limit_key=uid, source="env")


def _fake_agent(monkeypatch, seen):
    import agent as agent_mod

    def _answer(user_input, history=None, **kwargs):
        seen["input"] = str(user_input)
        return {"output": "ok", "active_tier": "Fake",
                "task_type": "research", "tools_used": [], "sources": []}

    monkeypatch.setattr(agent_mod, "answer_with_fallback", _answer)


def test_is_teaching_request():
    from backend.chatflow import _is_teaching_request

    assert _is_teaching_request("teach me these 2 slides one by one") is True
    assert _is_teaching_request("exam tomorrow, teach lecture-wise") is True
    assert _is_teaching_request("hello") is False
    assert _is_teaching_request("tere liye song") is False


def test_continuation_needs_active_session():
    from backend.chatflow import _is_teaching_continuation

    assert _is_teaching_continuation("Next", []) is False
    hist = [{"role": "assistant", "content": "📘 FILE: Lecture_01\nSlides: 1-3\nbla"}]
    assert _is_teaching_continuation("Next", hist) is True
    assert _is_teaching_continuation("next song", hist) is False
    assert _is_teaching_continuation("x" * 200, hist) is False


def test_legacy_header_still_parses():
    from backend.chatflow import _last_teaching_state, _match_teaching_header

    assert _match_teaching_header("📘 FILE: Deck\nSlides: 2-4") == ("Deck", 2, 4)
    assert _match_teaching_header("📘 FILE: Deck — Slides 2-4") == ("Deck", 2, 4)
    assert _match_teaching_header("no header") is None
    hist = [{"role": "assistant", "content": "📘 FILE: Deck — Slides 2-4\nbla"}]
    assert _last_teaching_state(hist) == ("Deck", 4)


def test_recall_answer_continues_session():
    from backend.chatflow import _is_teaching_continuation

    hist = [{"role": "assistant",
             "content": "📘 FILE: Lecture_01\nSlides: 1-3\n## Concept: Graph\n**Recall**\nWhat is a vertex?"}]
    assert _is_teaching_continuation("A vertex is a node", hist) is True
    assert _is_teaching_continuation("next song", hist) is False
    assert _is_teaching_continuation("y" * 400, hist) is False
    # No Recall checkpoint: plain answers do not continue.
    hist2 = [{"role": "assistant", "content": "📘 FILE: Lecture_01\nSlides: 1-3\nbla"}]
    assert _is_teaching_continuation("A vertex is a node", hist2) is False


def test_admin_block_detection():
    from backend.chatflow import _is_admin_block

    assert _is_admin_block("Course Code CSE 0613-4125 Credit 3.0 Instructor Rubel Sheikh Schedule Mondays") is True
    assert _is_admin_block("Attendance 10% Midterm 20% Final 40%") is True
    assert _is_admin_block("A graph is a pair (V, E) of vertices and edges") is False
    # Mixed definition + course code stays a concept.
    assert _is_admin_block("Graph theory studies graphs. Course CSE 0613-4125 covers the Handshaking theorem") is False


def test_concept_format_in_prompt_and_suffix():
    from agent.prompts import SYSTEM_PROMPT
    from backend.chatflow import TEACHING_SUFFIX

    for token in ("## Concept:", "**Definition**", "**Simple intuition**",
                  "**How it works**", "**Why it matters**", "**Example**",
                  "**Exam importance**", "**Exam trap**", "**Recall**", "**Source**"):
        assert token in SYSTEM_PROMPT
        assert token in TEACHING_SUFFIX
    assert "wait for the learner" in SYSTEM_PROMPT.lower()


def _norm_ws(text):
    return " ".join(str(text or "").split())


def test_texture_sentences_verbatim_in_both_prompts():
    from agent.prompts import SYSTEM_PROMPT
    from backend.chatflow import TEACHING_SUFFIX

    example_sentence = (
        "Never reuse the same conceptual example domain in consecutive turns. "
        "Rotate across genuinely different domains such as social networks → "
        "roads → circuits → food webs → databases, rather than merely changing "
        "names or surface details."
    )
    depth_sentence = (
        "omit any section that adds no meaningful information — do not "
        "artificially fill the canonical structure"
    )
    recall_sentence = "never ask the same recall type twice consecutively"
    for sentence in (example_sentence, depth_sentence, recall_sentence):
        assert _norm_ws(sentence) in _norm_ws(SYSTEM_PROMPT), sentence[:40]
        assert _norm_ws(sentence) in _norm_ws(TEACHING_SUFFIX), sentence[:40]


def test_reflection_teaching_focus():
    from agent.reflection import _TASK_FOCUS

    assert "source fidelity" in _TASK_FOCUS["research"]
    assert "source fidelity" in _TASK_FOCUS["multi_step"]


def test_router_teaching_never_simple():
    from agent.router import rule_route

    assert rule_route("teach me these 2 slides one by one") in ("research", "multi_step", "creative")
    assert rule_route("teach me these 2 slides one by one") != "simple"
    assert rule_route("exam tomorrow teach graphs") != "simple"
    # "example" must not be hijacked by the exam keyword (exam* stem bug).
    assert rule_route("give me an example sentence") is None


def test_extract_blocks_preserves_numbers(tmp_path, monkeypatch):
    from backend.chatflow import _extract_teaching_blocks

    ctx = _ctx(tmp_path, monkeypatch, "teach-blocks")
    b = _pptx_bytes([["Graph intro"], ["Walk definition"], ["Path definition"], ["Cycle"]])
    meta = ctx.file_store.save_upload(b, "Lecture_01_Graph_Introduction.pptx")
    blocks, total, status = _extract_teaching_blocks(
        ctx, {"id": meta.id, "kind": "document", "name": "Lecture_01_Graph_Introduction.pptx"})
    assert status == "OK"
    assert total == 4
    assert [n for n, _ in blocks] == [1, 2, 3, 4]
    assert "Walk" in blocks[1][1]


def test_teaching_picks_one_file_window(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-one")
    b1 = _pptx_bytes([["L1S1"], ["L1S2"], ["L1S3"], ["L1S4"], ["L1S5"]])
    b2 = _pptx_bytes([["L2S1"], ["L2S2"]])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01_Graph_Introduction.pptx")
    m2 = ctx.file_store.save_upload(b2, "Lecture_02_Practice_Problem_on_Lec_01.pptx")
    atts = [
        {"id": m1.id, "kind": "document", "name": "Lecture_01_Graph_Introduction.pptx"},
        {"id": m2.id, "kind": "document", "name": "Lecture_02_Practice_Problem_on_Lec_01.pptx"},
    ]
    send, vision, clarify = _apply_teaching_session(
        ctx, "teach me these 2 slides one by one", [], atts, [], "teach me")
    assert clarify is None
    # ONE file per batch: first sorted file's CONTENT only (analysis header
    # may name both files with counts, but must not leak L2 slide content).
    assert "Lecture_01" in send
    assert "L2S1" not in send
    assert "L1S1" in send and "L1S3" in send
    assert "L1S4" not in send  # window is 3 slides
    assert "Teaching mode" in send
    assert "Verified content" in send


def test_next_advances_window(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-next")
    b1 = _pptx_bytes([["L1S1"], ["L1S2"], ["L1S3"], ["L1S4"], ["L1S5"]])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01.pptx")
    hist = [
        {"role": "user", "content": "teach me", "attachments": [
            {"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]},
        {"role": "assistant", "content": "📘 FILE: Lecture_01.pptx\nSlides: 1-3\nbla"},
    ]
    send, _, _ = _apply_teaching_session(ctx, "Next", hist, [], [], "Next")
    assert "L1S4" in send and "L1S5" in send
    assert "L1S1" not in send


def test_exhausted_file_advances(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-adv")
    b1 = _pptx_bytes([["A1"], ["A2"]])
    b2 = _pptx_bytes([["B1"], ["B2"]])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01.pptx")
    m2 = ctx.file_store.save_upload(b2, "Lecture_02.pptx")
    hist = [
        {"role": "user", "content": "teach", "attachments": [
            {"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"},
            {"id": m2.id, "kind": "document", "name": "Lecture_02.pptx"}]},
        {"role": "assistant", "content": "📘 FILE: Lecture_01.pptx\nSlides: 1-2\n done."},
    ]
    send, _, _ = _apply_teaching_session(ctx, "Next", hist, [], [], "Next")
    assert "B1" in send  # advanced to second file
    assert "Teaching mode" in send


def test_no_docs_fail_closed(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-empty")
    send, _, clarify = _apply_teaching_session(ctx, "teach me slides", [], [], [], "teach me slides")
    assert clarify is None
    assert "no readable slides" in send.lower()
    assert "Do not invent" in send


def test_session_lost_files_short_circuits_no_model_call(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-lost")
    # Prior teaching exists, but no files are available anywhere now.
    hist = [
        {"role": "assistant", "content": "📘 FILE: Lecture_01.pptx\nSlides: 1-3\n## Concept: G\n**Recall**\nQ?"},
    ]
    send, _, clarify = _apply_teaching_session(ctx, "Next", hist, [], [], "Next")
    assert clarify is not None
    assert "re-upload" in clarify.lower()


def test_unreadable_window_short_circuits(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-broken")
    b1 = _pptx_bytes([["Graph vertex edge"]])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01.pptx")
    ctx.file_store.delete_upload(m1.id)
    atts = [{"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]
    send, _, clarify = _apply_teaching_session(
        ctx, "teach me slides", [], atts, [], "teach me slides")
    assert clarify is not None
    assert "re-upload" in clarify.lower()


def test_oversize_window_keeps_model_path(tmp_path, monkeypatch):
    import backend.chatflow as cf
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-big")
    monkeypatch.setattr(cf, "TEACHING_INLINE_MAX_BYTES", 10)
    b1 = _pptx_bytes([["Graph vertex edge"]])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01.pptx")
    atts = [{"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]
    send, _, clarify = _apply_teaching_session(
        ctx, "teach me slides", [], atts, [], "teach me slides")
    assert clarify is None
    assert "read_document" in send


def test_run_chat_teaching_injects_window(tmp_path, monkeypatch):
    from backend.chatflow import run_chat

    ctx = _ctx(tmp_path, monkeypatch, "teach-e2e")
    seen = {}
    _fake_agent(monkeypatch, seen)
    b1 = _pptx_bytes([["Graph intro vertex edge"], ["Walk repeats"], ["Path no repeat"], ["Cycle returns"]])
    b2 = _pptx_bytes([["Practice Q1"]])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01_Graph_Introduction.pptx")
    m2 = ctx.file_store.save_upload(b2, "Lecture_02_Practice_Problem_on_Lec_01.pptx")
    run_chat(ctx, "teach me these 2 slides one by one", upload_ids=[m1.id, m2.id])
    assert "Verified content" in seen["input"]
    assert "Teaching mode" in seen["input"]
    assert "Graph intro" in seen["input"]


def test_admin_compressed_in_analysis(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-admin")
    b1 = _pptx_bytes([
        ["Course Code CSE 0613-4125 Credit 3.0 Instructor Rubel Sheikh Schedule Mondays"],
        ["Attendance 10 percent Midterm 20 percent Final 40 percent grading policy"],
        ["A graph is a pair V E of vertices and edges"],
        ["Walk repeats vertices edges path cycle degree"],
    ])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01.pptx")
    atts = [{"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]
    send, _, _ = _apply_teaching_session(ctx, "teach me slides", [], atts, [], "teach me slides")
    assert "admin" in send.lower()
    assert "Concept:" in send  # new schema mirrored in suffix


def test_practice_quiz_revise_intents():
    from backend.chatflow import _is_teaching_request

    assert _is_teaching_request("quiz me on chapter 3") is True
    assert _is_teaching_request("give me practice problems on these slides") is True
    assert _is_teaching_request("revise lecture 2 notes") is True


def test_router_practice_quiz_research():
    from agent.router import rule_route

    assert rule_route("quiz me on this pdf") == "research"
    assert rule_route("give me practice questions") == "research"


def test_pace_feedback_detection():
    from backend.chatflow import _is_pace_feedback, _pace_direction

    hist = [{"role": "assistant", "content": "📘 FILE: L\nSlides: 1-3\n## Concept: G\n**Recall**\nQ?"}]
    assert _pace_direction("slow down, simpler please") == "slow"
    assert _pace_direction("got it, give harder problems") == "fast"
    assert _pace_direction("what is a vertex") is None
    assert _is_pace_feedback("slow down please", hist) is True
    assert _is_pace_feedback("slow down please", []) is False


def test_pace_note_injected(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-pace")
    b1 = _pptx_bytes([["Graph vertex edge"], ["Walk repeats"], ["Path simple"], ["Cycle closed"]])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01.pptx")
    hist = [
        {"role": "user", "content": "teach me", "attachments": [
            {"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]},
        {"role": "assistant", "content": "📘 FILE: Lecture_01.pptx\nSlides: 1-3\n## Concept: G\n**Recall**\nQ?"},
    ]
    send, _, _ = _apply_teaching_session(ctx, "slow down, simpler", hist, [], [], "slow down, simpler")
    assert "slow down" in send.lower()


def test_last_window_section_review(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-review")
    b1 = _pptx_bytes([["Graph vertex edge"], ["Walk repeats"]])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01.pptx")
    atts = [{"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]
    send, _, _ = _apply_teaching_session(ctx, "teach me slides", [], atts, [], "teach me slides")
    assert "section review" in send.lower()


def test_exhausted_all_files_exam_mode(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-exammode")
    b1 = _pptx_bytes([["Graph vertex edge"], ["Walk repeats"]])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01.pptx")
    hist = [
        {"role": "user", "content": "teach me", "attachments": [
            {"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]},
        {"role": "assistant", "content": "📘 FILE: Lecture_01.pptx\nSlides: 1-2\n## Concept: G\n**Recall**\nQ?"},
    ]
    send, _, _ = _apply_teaching_session(ctx, "Next", hist, [], [], "Next")
    assert "EXAM MODE" in send
    assert "Verified content" not in send


def test_single_block_exam_mode_carries_reference_text(tmp_path, monkeypatch):
    # Legacy single-block dumps (whole file == "slide 1"): cursor is past
    # the end on "next", so EXAM MODE fires with no fresh window. Citations
    # must still ground in the last taught source text, not memory.
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-examref")
    b1 = _pptx_bytes([["Osmosis moves water across membranes"]])
    m1 = ctx.file_store.save_upload(b1, "Solo_Lecture.pptx")
    hist = [
        {"role": "user", "content": "teach me", "attachments": [
            {"id": m1.id, "kind": "document", "name": "Solo_Lecture.pptx"}]},
        {"role": "assistant", "content": "📘 FILE: Solo_Lecture.pptx\nSlides: 1-1\n## Concept: G\n**Recall**\nQ?"},
    ]
    send, _, clarify = _apply_teaching_session(ctx, "Next", hist, [], [], "Next")
    assert clarify is None
    assert "EXAM MODE" in send
    assert "Verified content" not in send
    assert "Osmosis moves water across membranes" in send
    assert "[slide 1]" in send


def test_explicit_reteach_skips_exam_mode(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-reteach")
    b1 = _pptx_bytes([["Graph vertex edge"], ["Walk repeats"]])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01.pptx")
    hist = [
        {"role": "user", "content": "teach me", "attachments": [
            {"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]},
        {"role": "assistant", "content": "📘 FILE: Lecture_01.pptx\nSlides: 1-2\n## Concept: G\n**Recall**\nQ?"},
    ]
    send, _, _ = _apply_teaching_session(
        ctx, "teach Lecture_01 again", hist, [], [], "teach Lecture_01 again")
    assert "Verified content" in send


def test_time_pressure_detection():
    from backend.chatflow import _time_pressure

    assert _time_pressure("exam tomorrow, teach me quickly") == "rush"
    assert _time_pressure("teach me graphs in detail from scratch") == "deep"
    assert _time_pressure("teach me these slides") is None


def test_rush_modifier_injected(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-rush")
    b1 = _pptx_bytes([["Graph vertex edge"], ["Walk repeats"], ["Path simple"]])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01.pptx")
    atts = [{"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]
    send, _, _ = _apply_teaching_session(
        ctx, "exam tomorrow teach quickly", [], atts, [], "exam tomorrow teach quickly")
    assert "essentials" in send.lower()


def test_format_logging_metadata_only(tmp_path, monkeypatch):
    import backend.chatflow as cf
    import backend.teach as teach_mod

    seen = {}
    monkeypatch.setattr(teach_mod, "obs_event", lambda name, **kw: seen.update({"name": name, **kw}))
    cf._log_teaching_format(
        "hi [Teaching mode: x]",
        "📘 FILE: L\nSlides: 1-2\n## Concept: G\n**Source**\n[slide 1]\n**Recall**\nQ?",
        "Cohere")
    assert seen["name"] == "teaching.format"
    assert seen["has_concept"] is True and seen["has_recall"] is True
    assert seen["tier"] == "Cohere"
    # Non-teaching turns log nothing.
    seen.clear()
    cf._log_teaching_format("hello", "hi there", "Groq")
    assert seen == {}
    # Never raises, even on hostile input.
    cf._log_teaching_format(None, None, None)


def test_texture_logging_fields(tmp_path, monkeypatch):
    import backend.chatflow as cf
    import backend.teach as teach_mod

    seen = {}
    monkeypatch.setattr(teach_mod, "obs_event", lambda name, **kw: seen.update({"name": name, **kw}))
    cf._log_teaching_format(
        "hi [Teaching mode: x]",
        "📘 FILE: L\nSlides: 1-2\n## Concept: G\n**Definition**\nA pair.\n"
        "**How it works**\nLinks.\n**Exam importance**\nMEDIUM stuff.\n"
        "**Source**\n[slide 1]\n**Recall**\nWhy does this hold?",
        "Groq")
    assert seen["has_definition"] is True
    assert seen["has_how"] is True
    assert seen["has_why"] is False
    assert seen["has_example"] is False
    assert seen["importance"] == "MEDIUM"
    assert seen["recall_type"] == "why"


def test_classify_recall_type():
    from backend.chatflow import _classify_recall_type

    assert _classify_recall_type("**Recall**\nWhat is a vertex?") == "define"
    assert _classify_recall_type("**Recall**\nCompare paths and cycles") == "compare"
    assert _classify_recall_type("**Recall**\nFind the mistake: ...") == "mistake"
    assert _classify_recall_type("**Recall**\nSolve for x") == "apply"
    assert _classify_recall_type("no recall here") == "unknown"
    assert _classify_recall_type(None) == "unknown"


_GOOD_DRAFT = (
    "📘 FILE: Lecture_01.pptx\nSlides: 4-6\n"
    "## Concept: Graph\n**Definition**\nA pair (V, E).\n"
    "**Source**\n[slide 4]\n**Recall**\nWhat is V?"
)


def test_scope_fence_names_window():
    from backend.chatflow import _teaching_scope_from_send, _teaching_scope_line

    line = _teaching_scope_line(4, 6, 44)
    assert "ONLY slides 4-6" in line
    assert "7+ are NOT loaded" in line
    assert _teaching_scope_from_send("x\n" + line) == (4, 6)
    assert _teaching_scope_from_send("no fence here") is None


def test_validator_rejects_fabrication_admissions():
    from backend.chatflow import _validate_teaching_draft

    base = _GOOD_DRAFT
    assert any("unverified" in r for r in _validate_teaching_draft(
        base + "\n(Assumes Lecture 5 slides are as follows — adjust if actual slides differ.)", 4, 6))
    assert any("unverified" in r for r in _validate_teaching_draft(
        base + "\nAlways verify slides before teaching!", 4, 6))
    # Legitimate prose about assumptions must not trip the detector.
    assert _validate_teaching_draft(
        base + "\nAssume costs are non-negative for this proof.", 4, 6) == []


def test_validator_rejects_slide_dump():
    from backend.chatflow import _validate_teaching_draft

    dump = ("📘 FILE: L\nSlides: 4-6\n## Concept: Big\n**Source**\n[slide 4]\n"
            "**Source**\n[slide 5]\n**Source**\n[slide 6]\n**Source**\n[slide 7]\n"
            "**Recall**\nQ?")
    reasons = _validate_teaching_draft(dump, 4, 6)
    assert any("beyond window" in r for r in reasons)
    dump_inside = ("📘 FILE: L\nSlides: 4-6\n## Concept: Big\n**Source**\n[slide 4]\n"
                   "**Source**\n[slide 5]\n**Source**\n[slide 6]\n**Recall**\nQ?")
    assert _validate_teaching_draft(dump_inside, 4, 6) == []
    # Backward references to earlier slides stay allowed (bridges, not dumps).
    bridged = dump_inside.replace("[slide 6]", "[slide 2]")
    assert _validate_teaching_draft(bridged, 4, 6) == []


def test_validator_matrix():
    from backend.chatflow import _validate_teaching_draft

    assert _validate_teaching_draft(_GOOD_DRAFT, 4, 6) == []
    assert _validate_teaching_draft("", 4, 6) == ["empty answer"]
    overflow = _GOOD_DRAFT.replace("Slides: 4-6", "Slides: 4-10")
    assert any("outside window" in r for r in _validate_teaching_draft(overflow, 4, 6))
    future = _GOOD_DRAFT + "\n[slide 9]"
    assert any("beyond window" in r for r in _validate_teaching_draft(future, 4, 6))
    rangecite = _GOOD_DRAFT + "\n**Source**\n[slides 4-9]"
    assert any("beyond window" in r for r in _validate_teaching_draft(rangecite, 4, 6))
    nocite = ("📘 FILE: L\nSlides: 4-6\n## Concept: G\n**Recall**\nQ?")
    assert any("no slide citations" in r for r in _validate_teaching_draft(nocite, 4, 6))
    foot = _GOOD_DRAFT + "\nSay Next for slides 7-9."
    assert any("banned footer" in r for r in _validate_teaching_draft(foot, 4, 6))
    gotit = _GOOD_DRAFT + '\nSay "Got it" when ready.'
    assert any("banned footer" in r for r in _validate_teaching_draft(gotit, 4, 6))
    noconcept = ("📘 FILE: L\nSlides: 4-6\nSome table here.\n**Source**\n[slide 4]\n"
                 "**Recall**\nQ?")
    assert any("no Concept block" in r for r in _validate_teaching_draft(noconcept, 4, 6))
    tworecall = _GOOD_DRAFT + "\n**Recall**\nAnother?"
    assert any("Recall sections" in r for r in _validate_teaching_draft(tworecall, 4, 6))
    earlyrecall = ("📘 FILE: L\nSlides: 4-6\n**Recall**\nQ?\n## Concept: G\n"
                   "**Source**\n[slide 4]")
    assert any("after the last Concept" in r for r in _validate_teaching_draft(earlyrecall, 4, 6))
    noheader = ("## Concept: G\n**Source**\n[slide 4]\n**Recall**\nQ?")
    assert any("FILE header" in r for r in _validate_teaching_draft(noheader, 4, 6))
    admin = ("📘 FILE: L\nSlides: 1-3\n### Administrative Information\n"
             "- Course Code 0613-4125\n**Source:** [slide 2]")
    assert _validate_teaching_draft(admin, 1, 3) == []
    adminrecall = admin + "\n**Recall**\nQ?"
    assert any("admin-only" in r for r in _validate_teaching_draft(adminrecall, 1, 3))


def test_paced_turn_with_continue_cue_passes():
    """Word-budget pacing: a short turn ending in the continue cue is valid.

    The `Reply **continue**` navigation line must not trip the banned
    footer detector (only Say Next / Say Got it / Next Steps are banned),
    and the paced shape (header + concept + one terminal Recall) passes.
    """
    from backend.chatflow import _validate_teaching_draft

    paced = (_GOOD_DRAFT + "\nReply **continue** for the next concept.")
    assert _validate_teaching_draft(paced, 4, 6) == []


def test_pacing_rule_in_prompt_and_suffix():
    """Word-budget rule + continue cue live in both teaching contracts."""
    from agent.prompts import SYSTEM_PROMPT
    from backend.chatflow import TEACHING_SUFFIX

    for token in ("~300 words", "Reply **continue** for the next concept.",
                  "Never re-teach"):
        assert token in SYSTEM_PROMPT, token
        assert token in TEACHING_SUFFIX, token


def test_recall_without_answer_key_in_prompt_and_suffix():
    """Recall asks only; logistics are never invented (both contracts)."""
    from agent.prompts import SYSTEM_PROMPT
    from backend.chatflow import TEACHING_SUFFIX

    for token in ("never its answer or answer key",
                  "Never invent test dates"):
        assert token in SYSTEM_PROMPT, token
        assert token in TEACHING_SUFFIX, token


def test_duplicate_file_header_fails_validation():
    """Two source headers trip validation so repair can fix them."""
    from backend.chatflow import _validate_teaching_draft

    doubled = _GOOD_DRAFT + "\n📘 FILE: Other.pptx\nSlides: 4-6"
    assert any("duplicate FILE header" in r
               for r in _validate_teaching_draft(doubled, 4, 6))
    assert _validate_teaching_draft(_GOOD_DRAFT, 4, 6) == []


def test_upload_id_redaction():
    """Echoed upload IDs are withheld; download IDs keep working."""
    from backend.teach import _redact_upload_ids

    leaked = "Source of the file (upload ID: 12715c1de4374507)."
    fixed = _redact_upload_ids(leaked)
    assert "12715c1de4374507" not in fixed
    assert "upload ID" in fixed
    delivery = "Presentation saved as deck.pptx (file ID: a3b5ecfce1df4287)"
    assert _redact_upload_ids(delivery) == delivery
    # Longer hashes (sessions, SHAs) are not partial-matched.
    long_hash = "upload ID: " + "ab12" * 16
    assert _redact_upload_ids(long_hash) == long_hash
    assert _redact_upload_ids(None) is None
    assert _redact_upload_ids(123) == 123


def test_repair_turn_redacts_upload_ids():
    """The persisted teaching answer never carries an upload ID."""
    from backend.teach import _maybe_repair_teaching_turn

    send = "x teach ONLY slides 1-1 y"
    content = ("📘 FILE: L.pptx\nSlides: 1-1\n## Concept: G\n"
               "**Source**\n[slide 1] (upload ID: 12715c1de4374507)\n"
               "**Recall**\nQ?")
    fixed, _repaired, _left = _maybe_repair_teaching_turn(send, content, "Groq")
    assert "12715c1de4374507" not in fixed
    assert "[slide 1]" in fixed


def _stub_repair(monkeypatch, text):
    import types

    import agent as agent_mod
    import config

    calls = {}

    def fake_llm(name, temperature=0.3):
        calls["tier"] = name
        return object()

    def fake_invoke(llm, messages, **kw):
        calls["kw"] = kw
        return types.SimpleNamespace(content=text)

    monkeypatch.setattr(config, "get_tier_llm", fake_llm)
    monkeypatch.setattr(agent_mod, "_invoke_bounded", fake_invoke)
    return calls


def test_repair_fixes_and_reports(monkeypatch):
    from backend.chatflow import _maybe_repair_teaching_turn

    send = "window\n" + "x teach ONLY slides 4-6 y"
    bad = "Concept: G\nSource: [slide 4]\nSource: [slide 9]"
    calls = _stub_repair(monkeypatch, _GOOD_DRAFT)
    fixed, repaired, left = _maybe_repair_teaching_turn(send, bad, "Cohere")
    assert repaired is True and fixed == _GOOD_DRAFT and left == []
    # Cross-tier repair: drafts are fixed on the strongest live
    # tier (Groq first), not the failed tier itself.
    assert calls["tier"] == "Groq"


def test_repair_keeps_draft_when_unfixable(monkeypatch):
    import agent as agent_mod
    import config
    from backend.chatflow import _maybe_repair_teaching_turn

    monkeypatch.setattr(config, "get_tier_llm", lambda name, temperature=0.3: None)
    called = []
    monkeypatch.setattr(agent_mod, "_invoke_bounded",
                        lambda *a, **k: called.append(1) or None)
    send = "window\n" + "x teach ONLY slides 4-6 y"
    bad = "no header here"
    fixed, repaired, left = _maybe_repair_teaching_turn(send, bad, "Nope")
    assert repaired is False and fixed == bad and called == [] and left


def test_repair_resets_stream(monkeypatch):
    from backend.chatflow import _maybe_repair_teaching_turn

    send = "window\n" + "x teach ONLY slides 4-6 y"
    bad = "Concept: G\nSource: [slide 9]"
    _stub_repair(monkeypatch, _GOOD_DRAFT)
    resets = []
    fixed, repaired, _ = _maybe_repair_teaching_turn(
        send, bad, "Groq", on_reset=lambda: resets.append(1))
    assert repaired is True and resets == [1] and fixed == _GOOD_DRAFT


def test_maybe_repair_skips_non_teaching(monkeypatch):
    import agent as agent_mod
    from backend.chatflow import _maybe_repair_teaching_turn

    calls = []
    monkeypatch.setattr(agent_mod, "_invoke_bounded", lambda *a, **k: calls.append(1))
    fixed, repaired, left = _maybe_repair_teaching_turn("hello", "hi", "Groq")
    assert (fixed, repaired, left) == ("hi", False, [])
    assert calls == []


def test_scope_fence_in_session_output(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-scope")
    b1 = _pptx_bytes([["Graph vertex edge"], ["Walk repeats"], ["Path simple"]])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01.pptx")
    atts = [{"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]
    send, _, clarify = _apply_teaching_session(
        ctx, "teach me slides", [], atts, [], "teach me slides")
    assert clarify is None
    assert "Scope fence" in send and "ONLY slides 1-3" in send


def test_subject_templates_in_prompt():
    from agent.prompts import SYSTEM_PROMPT

    assert "line-by-line" in SYSTEM_PROMPT
    assert "compact comparison" in SYSTEM_PROMPT
    assert "missing prerequisite first" in SYSTEM_PROMPT


def test_recall_answer_gets_evaluation_note(tmp_path, monkeypatch):
    from backend.chatflow import _apply_teaching_session

    ctx = _ctx(tmp_path, monkeypatch, "teach-eval")
    b1 = _pptx_bytes([["Graph vertex edge"], ["Walk repeats"], ["Path simple"], ["Cycle closed"]])
    m1 = ctx.file_store.save_upload(b1, "Lecture_01.pptx")
    hist = [
        {"role": "user", "content": "teach me", "attachments": [
            {"id": m1.id, "kind": "document", "name": "Lecture_01.pptx"}]},
        {"role": "assistant", "content": "📘 FILE: Lecture_01.pptx\nSlides: 1-3\n## Concept: Graph\n**Recall**\nWhat is a vertex?"},
    ]
    send, _, _ = _apply_teaching_session(ctx, "A vertex is a node", hist, [], [], "A vertex is a node")
    assert "evaluate" in send.lower()
    assert "Recall checkpoint" in send


def test_repair_bills_parent_budget(monkeypatch):
    """Repair charges the passed request budget (no unbilled side budget)."""
    import types

    import agent as agent_mod
    from agent.budget import RequestBudget
    from backend.chatflow import _maybe_repair_teaching_turn

    import config

    seen = {}

    def fake_llm(name, temperature=0.3):
        return object()

    def fake_invoke(llm, messages, budget=None, **kw):
        seen["budget"] = budget
        if budget is not None:
            budget.count_llm()
        return types.SimpleNamespace(content=_GOOD_DRAFT)

    monkeypatch.setattr(config, "get_tier_llm", fake_llm)
    monkeypatch.setattr(agent_mod, "_invoke_bounded", fake_invoke)
    parent = RequestBudget()
    send = "window\n" + "x teach ONLY slides 4-6 y"
    bad = "Concept: G\nSource: [slide 4]\nSource: [slide 9]"
    fixed, repaired, left = _maybe_repair_teaching_turn(
        send, bad, "Groq", budget=parent)
    assert repaired is True and fixed == _GOOD_DRAFT and left == []
    assert seen["budget"] is parent
    assert parent.llm_calls == 1


def test_repair_on_spent_budget_keeps_draft(monkeypatch):
    """Exhausted parent budget fails fast into the untouched draft."""
    import types

    import agent as agent_mod
    from agent.budget import RequestBudget
    from backend.chatflow import _maybe_repair_teaching_turn

    import config

    calls = []

    def fake_invoke(llm, messages, budget=None, **kw):
        calls.append(1)
        if budget is not None:
            budget.count_llm()
        return types.SimpleNamespace(content=_GOOD_DRAFT)

    monkeypatch.setattr(config, "get_tier_llm", lambda name, temperature=0.3: object())
    monkeypatch.setattr(agent_mod, "_invoke_bounded", fake_invoke)
    spent = RequestBudget(max_llm=0)
    send = "window\n" + "x teach ONLY slides 4-6 y"
    bad = "Concept: G\nSource: [slide 4]\nSource: [slide 9]"
    fixed, repaired, left = _maybe_repair_teaching_turn(
        send, bad, "Groq", budget=spent)
    assert repaired is False and fixed == bad and left
    # One attempt charged against the spent budget, then fail-fast into
    # the untouched draft (no unbilled side budget, no partial repair).
    assert calls == [1]
