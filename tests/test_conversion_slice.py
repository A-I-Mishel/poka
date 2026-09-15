"""Conversion slice: structured DOCX read + Word/doc gate + read->create path."""

import contextlib
import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from services import context as ctx


@contextlib.contextmanager
def _user(tmp_path, monkeypatch, uid="conv-user"):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("PLUTO_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("PLUTO_USER_ID", uid)
    monkeypatch.delenv("PLUTO_AUTH_MODE", raising=False)
    ctx.set_current_user_id(uid)
    from backend.deps import UserContext
    from services.files import FileStore
    from services.storage import UserStore

    try:
        yield UserContext(user_id=uid, user_store=UserStore(uid),
                          file_store=FileStore(uid), limit_key=uid, source="env")
    finally:
        ctx.set_current_user_id(None)


def _docx_bytes():
    from docx import Document

    doc = Document()
    doc.add_heading("Report Title", level=1)
    doc.add_heading("Findings", level=2)
    doc.add_paragraph("Intro paragraph here.")
    doc.add_paragraph("First point", style="List Bullet")
    doc.add_paragraph("Step one", style="List Number")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Asha"
    table.cell(1, 1).text = "95"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_read_preserves_structure(tmp_path, monkeypatch):
    from tools.document_tool import read_document

    with _user(tmp_path, monkeypatch) as uctx:
        meta = uctx.file_store.save_upload(_docx_bytes(), "report.docx")
        assert meta.kind == "document"
        text = read_document.invoke({"upload_id": meta.id})
    assert not text.startswith("STATUS=")
    assert "# Report Title" in text
    assert "## Findings" in text
    assert "- First point" in text
    assert "1. Step one" in text
    assert "Name | Score" in text and "Asha | 95" in text


def test_gate_bare_word_request_uses_doc():
    from agent.attachment_gate import decide

    docs = [{"id": "d" * 16, "kind": "document", "name": "report.docx"}]
    for phrase in ("Convert the Word file.", "Convert this doc."):
        d = decide(phrase, [], docs)
        assert [e["id"] for e in d["use_docs"]] == ["d" * 16], (phrase, d)
        assert d["clarify"] is None


def test_run_chat_convert_word_to_pdf(tmp_path, monkeypatch):
    import agent as agent_mod
    from backend.chatflow import run_chat

    with _user(tmp_path, monkeypatch, "conv-chat") as uctx:
        meta = uctx.file_store.save_upload(_docx_bytes(), "report.docx")
        seen = {}

        def _answer(user_input, history=None, **kwargs):
            seen["image_ids"] = list(kwargs.get("image_upload_ids") or [])
            seen["input"] = str(user_input)
            return {"output": "ok", "active_tier": "Fake",
                    "task_type": "simple", "tools_used": [], "sources": []}

        monkeypatch.setattr(agent_mod, "answer_with_fallback", _answer)
        run_chat(uctx, "(attachment)", upload_ids=[meta.id])
        run_chat(uctx, "Convert this Word file into PDF.")
        assert seen["image_ids"] == []
        assert meta.id in seen["input"]


def test_read_then_create_pdf(tmp_path, monkeypatch):
    from tools.document_tool import read_document
    from tools.make_tool import create_pdf

    with _user(tmp_path, monkeypatch, "conv-pdf") as uctx:
        meta = uctx.file_store.save_upload(_docx_bytes(), "report.docx")
        text = read_document.invoke({"upload_id": meta.id})
        assert not text.startswith("STATUS=")
        out = create_pdf.invoke({"title": "Report", "markdown_text": text})
        assert "file ID:" in out
        outs = uctx.file_store.list_outputs()
        assert outs and outs[0].kind == "pdf"
